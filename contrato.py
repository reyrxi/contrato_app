import json
import os
import threading
import tkinter as tk
from tkinter import filedialog, messagebox, ttk

from docx import Document

from constants import (
    ACCENT, ACCENT_L, BG, BORDER, BORDER2, CARD,
    DANGER, DRAFT_FILE, F10, F10B, F13B, F8, F9, F9B, GUIDE,
    SIDE_HOV, SIDE_SEL, SIDE_TXT, SIDEBAR, SUCCESS, TEXT, TEXT2, TEXT3, WARN, WHITE,
)
from document import convert_to_pdf, replace_all
from utils import HAS_OPENPYXL, HAS_REQUESTS, buscar_cep, fmt_cur
from widgets import Btn, LabeledEntry, ScrollableFrame, Section, build_fields


class ContratoApp(tk.Tk):
    PAGES = [
        ("Contratante", "👤", "Dados do contratante"),
        ("Curso",       "📚", "Curso, datas e valores"),
        ("Responsável", "💳", "Responsável financeiro"),
        ("Cursos",      "📋", "Gerenciar cursos"),
        ("Config",      "⚙",  "Configurações"),
        ("Placeholders","📌", "Guia de placeholders"),
    ]

    def __init__(self):
        super().__init__()
        self.title("Contratos de Matrícula")
        self.geometry("1060x720")
        self.minsize(860, 620)
        self.configure(bg=BG)

        self.v = {}
        self.cursos_lista = []
        self._resp_entries = []
        self.v["_modelo_path"] = tk.StringVar(value="")
        self._cur = ""

        self._build_ui()
        self._load_draft()

    def _build_ui(self):
        self._sidebar = tk.Frame(self, bg=SIDEBAR, width=220)
        self._sidebar.pack(side="left", fill="y")
        self._sidebar.pack_propagate(False)
        self._build_sidebar()

        right = tk.Frame(self, bg=BG)
        right.pack(side="left", fill="both", expand=True)

        self._topbar = tk.Frame(
            right, bg=WHITE, highlightbackground=BORDER, highlightthickness=1, height=60
        )
        self._topbar.pack(fill="x")
        self._topbar.pack_propagate(False)
        self._build_topbar()

        self._body = tk.Frame(right, bg=BG)
        self._body.pack(fill="both", expand=True)

        statusbar = tk.Frame(
            right, bg=WHITE, highlightbackground=BORDER, highlightthickness=1, height=28
        )
        statusbar.pack(fill="x", side="bottom")
        statusbar.pack_propagate(False)
        self._status_lbl = tk.Label(
            statusbar, text="Pronto", font=F9, bg=WHITE, fg=TEXT3, anchor="w"
        )
        self._status_lbl.pack(side="left", padx=16, pady=4)
        tk.Label(statusbar, text="v3.0", font=F9, bg=WHITE, fg=TEXT3).pack(
            side="right", padx=16
        )

        self._pages = {
            "Contratante":  self._pg_contratante(),
            "Curso":        self._pg_curso(),
            "Responsável":  self._pg_resp(),
            "Cursos":       self._pg_cursos(),
            "Config":       self._pg_config(),
            "Placeholders": self._pg_placeholders(),
        }
        self._goto("Contratante")

    def _build_sidebar(self):
        logo = tk.Frame(self._sidebar, bg=SIDEBAR, pady=24)
        logo.pack(fill="x")
        tk.Label(logo, text="📝", font=("Segoe UI", 30), bg=SIDEBAR, fg=WHITE).pack()
        tk.Label(logo, text="Contratos", font=("Segoe UI", 13, "bold"), bg=SIDEBAR, fg=WHITE).pack()
        tk.Label(logo, text="de Matrícula", font=("Segoe UI", 10), bg=SIDEBAR, fg=SIDE_TXT).pack()
        tk.Frame(self._sidebar, bg=SIDE_SEL, height=1).pack(fill="x", padx=20, pady=(0, 8))

        self._nav = {}
        for name, icon, desc in self.PAGES:
            frame = tk.Frame(self._sidebar, bg=SIDEBAR, cursor="hand2")
            frame.pack(fill="x", padx=10, pady=2)
            inner = tk.Frame(frame, bg=SIDEBAR)
            inner.pack(fill="x", padx=6, pady=8)
            tk.Label(inner, text=icon, font=("Segoe UI", 13), bg=SIDEBAR, fg=SIDE_TXT, width=2).pack(
                side="left"
            )
            text_frame = tk.Frame(inner, bg=SIDEBAR)
            text_frame.pack(side="left", padx=8)
            tk.Label(text_frame, text=name, font=F10B, bg=SIDEBAR, fg=SIDE_TXT, anchor="w").pack(
                anchor="w"
            )
            tk.Label(text_frame, text=desc, font=F8, bg="#475569", anchor="w").pack(anchor="w")

            self._nav[name] = (frame, inner, text_frame)
            all_widgets = (
                [frame, inner, text_frame]
                + list(inner.winfo_children())
                + list(text_frame.winfo_children())
            )
            for w in all_widgets:
                w.bind("<Button-1>", lambda e, n=name: self._goto(n))
                w.bind("<Enter>", lambda e, ws=all_widgets: self._set_bg(ws, SIDE_HOV))
                w.bind(
                    "<Leave>",
                    lambda e, ws=all_widgets, n=name: self._set_bg(
                        ws, SIDE_SEL if n == self._cur else SIDEBAR
                    ),
                )
        tk.Label(self._sidebar, text="v3.0", font=F8, bg=SIDEBAR, fg="#334155").pack(
            side="bottom", pady=8
        )

    def _build_topbar(self):
        left = tk.Frame(self._topbar, bg=WHITE)
        left.pack(side="left", padx=20, pady=10)
        self._page_icon = tk.Label(left, text="👤", font=("Segoe UI", 16), bg=WHITE, fg=ACCENT)
        self._page_icon.pack(side="left", padx=(0, 8))
        title_frame = tk.Frame(left, bg=WHITE)
        title_frame.pack(side="left")
        self._page_title = tk.Label(title_frame, text="Contratante", font=F13B, bg=WHITE, fg=TEXT)
        self._page_title.pack(anchor="w")
        self._page_sub = tk.Label(
            title_frame, text="Dados do contratante", font=F9, bg=WHITE, fg=TEXT2
        )
        self._page_sub.pack(anchor="w")

        right = tk.Frame(self._topbar, bg=WHITE)
        right.pack(side="right", padx=16, pady=12)
        Btn(right, "⚡ Gerar Ambos", lambda: self.gerar("ambos"), SUCCESS).pack(
            side="right", padx=(8, 0)
        )
        Btn(right, "📕 PDF",  lambda: self.gerar("pdf"),   ACCENT, light=True).pack(
            side="right", padx=4
        )
        Btn(right, "📄 Word", lambda: self.gerar("word"),  ACCENT, light=True).pack(
            side="right", padx=4
        )
        tk.Frame(right, bg=BORDER, width=1).pack(side="right", fill="y", padx=8)
        Btn(right, "💾 Salvar",  self._save_draft,     "#0EA5E9", small=True).pack(
            side="right", padx=4
        )
        Btn(right, "🗑 Limpar", self._confirm_clear, "#94A3B8", small=True).pack(
            side="right", padx=4
        )

    def _goto(self, name):
        self._cur = name
        page_info = {p[0]: (p[1], p[2]) for p in self.PAGES}
        for n, (frame, inner, text_frame) in self._nav.items():
            is_selected = n == name
            bg = SIDE_SEL if is_selected else SIDEBAR
            all_widgets = (
                [frame, inner, text_frame]
                + list(inner.winfo_children())
                + list(text_frame.winfo_children())
            )
            self._set_bg(all_widgets, bg)
        if name in page_info:
            icon, sub = page_info[name]
            self._page_icon.configure(text=icon)
            self._page_title.configure(text=name)
            self._page_sub.configure(text=sub)
        for page in self._pages.values():
            page.pack_forget()
        self._pages[name].pack(fill="both", expand=True)

    def _set_bg(self, widgets, bg):
        for w in widgets:
            try:
                w.configure(bg=bg)
            except tk.TclError:
                pass

    def sv(self, key, default=""):
        if key not in self.v:
            self.v[key] = tk.StringVar(value=default)
        return self.v[key]

    def _status(self, msg, color=TEXT2):
        self._status_lbl.configure(text=msg, fg=color)

    def _pg_contratante(self):
        sc = ScrollableFrame(self._body)
        f = sc.inner

        s1 = Section(f, "Identificação", "👤")
        s1.pack(fill="x", padx=20, pady=(16, 6))
        build_fields(s1, [
            ("Nome completo",  "nome",         30, "text"),
            ("Naturalidade",   "naturalidade", 22, "text"),
            ("Data de nasc.",  "data_nasc",    14, "text", "DD/MM/AAAA"),
            ("RG",             "rg",           18, "text"),
            ("CPF",            "cpf",          18, "cpf",  "000.000.000-00"),
            ("E-mail",         "email",        28, "text"),
        ], self.v, self)

        s2 = Section(f, "Endereço — CEP preenche automaticamente", "📍")
        s2.pack(fill="x", padx=20, pady=6)
        build_fields(s2, [
            ("CEP",         "cep",         10, "cep",  "00000-000"),
            ("Número",      "numero",       8, "text"),
            ("Logradouro",  "endereco",    30, "text"),
            ("Complemento", "complemento", 20, "text"),
            ("Bairro",      "bairro",      22, "text"),
            ("Cidade",      "cidade",      22, "text"),
            ("UF",          "uf",           5, "text"),
        ], self.v, self)

        s3 = Section(f, "Contato", "📱")
        s3.pack(fill="x", padx=20, pady=(6, 20))
        build_fields(s3, [
            ("Telefone 1", "telefone1", 16, "fone", "(XX) XXXXX-XXXX"),
            ("Telefone 2", "telefone2", 16, "fone", "(XX) XXXXX-XXXX"),
        ], self.v, self)

        return sc

    def _pg_curso(self):
        sc = ScrollableFrame(self._body)
        f = sc.inner

        sc1 = Section(f, "Curso Escolhido", "📚")
        sc1.pack(fill="x", padx=20, pady=(16, 6))

        combo_frame = tk.Frame(sc1, bg=CARD)
        combo_frame.pack(fill="x", padx=20, pady=(8, 4))
        tk.Label(combo_frame, text="Curso", font=F9B, bg=CARD, fg=TEXT2).pack(anchor="w")
        row = tk.Frame(combo_frame, bg=CARD)
        row.pack(fill="x", pady=(3, 0))
        self.v["curso"] = tk.StringVar()
        style = ttk.Style()
        style.configure("C.TCombobox", fieldbackground=WHITE, background=WHITE, foreground=TEXT, font=F10)
        self._curso_combo = ttk.Combobox(
            row, textvariable=self.v["curso"], font=F10, width=40, style="C.TCombobox"
        )
        self._curso_combo.pack(side="left", ipady=5)
        self._curso_combo.bind("<<ComboboxSelected>>", self._on_combo_curso)
        Btn(row, "➕ Gerenciar", lambda: self._goto("Cursos"), ACCENT, light=True, small=True).pack(
            side="left", padx=10
        )
        tk.Label(
            sc1,
            text="  💡 Horário preenche automaticamente ao selecionar o curso.",
            font=F9, bg=CARD, fg=TEXT3,
        ).pack(anchor="w", padx=20, pady=(4, 12))
        build_fields(sc1, [
            ("Dias e Horário", "dias_horario", 36, "text", "Ex: Seg/Qua/Sex 19h–22h"),
        ], self.v, self)

        s2 = Section(f, "Datas de Previsão de Início", "📅")
        s2.pack(fill="x", padx=20, pady=6)
        build_fields(s2, [
            ("Opção 1", "inicio_1", 14, "text", "DD/MM/AAAA"),
            ("Opção 2", "inicio_2", 14, "text", "DD/MM/AAAA"),
            ("Opção 3", "inicio_3", 14, "text", "DD/MM/AAAA"),
        ], self.v, self)

        s3 = Section(f, "Valores Financeiros", "💰")
        s3.pack(fill="x", padx=20, pady=6)
        build_fields(s3, [
            ("Valor Total do Curso",     "valor_total",        14, "text", "R$"),
            ("Parcela Normal",           "parcela_normal",     14, "text", "R$"),
            ("Desconto (%) até dia 10",  "desconto_pct",        8, "text", "%"),
            ("Parcela c/ Desconto",      "parcela_desc",       14, "text", "R$ (auto)"),
            ("Quantidade de Parcelas",   "qtd_parcelas",        8, "text"),
            ("Dia de Vencimento",        "dia_vencimento",      8, "text", "dia"),
            ("Valor da Matrícula",       "valor_matricula",    14, "text", "R$"),
            ("Data do 1º Pagamento",     "data_1pag",          14, "text"),
            ("Mês de Ref. 1ª Parcela",   "mes_ref",            16, "text"),
            ("Data Venc. s/ Desconto",   "data_venc_sem_desc", 14, "text"),
            ("Forma de Pagamento",       "forma_pagamento",    22, "text"),
        ], self.v, self)
        calc_frame = tk.Frame(s3, bg=CARD)
        calc_frame.pack(fill="x", padx=20, pady=(0, 16))
        Btn(
            calc_frame, "🧮 Calcular Parcela com Desconto", self._calc_desc, ACCENT, light=True, small=True
        ).pack(anchor="w")

        s4 = Section(f, "Observações", "📝")
        s4.pack(fill="x", padx=20, pady=(6, 20))
        obs_frame = tk.Frame(s4, bg=CARD)
        obs_frame.pack(fill="x", padx=20, pady=(8, 16))
        tk.Label(obs_frame, text="Observações", font=F9B, bg=CARD, fg=TEXT2).pack(anchor="w")
        border = tk.Frame(obs_frame, bg=BORDER2)
        border.pack(fill="x", pady=(3, 0))
        inner = tk.Frame(border, bg=WHITE)
        inner.pack(fill="x", padx=1, pady=1)
        self._obs = tk.Text(
            inner, height=4, font=F10, bg=WHITE, fg=TEXT, bd=0, relief="flat",
            wrap="word", insertbackground=ACCENT,
        )
        self._obs.pack(fill="x", padx=10, pady=8)

        return sc

    def _calc_desc(self):
        try:
            normal = float(
                self.v["parcela_normal"].get().replace(",", ".").replace("R$", "").strip()
            )
            pct = float(
                self.v["desconto_pct"].get().replace(",", ".").replace("%", "").strip() or "0"
            )
            self.v["parcela_desc"].set(f"{normal * (1 - pct / 100):.2f}".replace(".", ","))
            self._status("✓ Calculado!", SUCCESS)
        except ValueError:
            messagebox.showwarning("Atenção", "Preencha Parcela Normal e Desconto corretamente.")

    def _pg_resp(self):
        sc = ScrollableFrame(self._body)
        f = sc.inner

        chk_sec = Section(f, "", "")
        chk_sec.pack(fill="x", padx=20, pady=(16, 4))
        cr = tk.Frame(chk_sec, bg=CARD)
        cr.pack(fill="x", padx=20, pady=14)
        self.v["_mesmo"] = tk.BooleanVar(value=False)

        toggle = tk.Frame(cr, bg=ACCENT_L, cursor="hand2", highlightbackground=ACCENT, highlightthickness=1)
        toggle.pack(fill="x")
        ci = tk.Frame(toggle, bg=ACCENT_L)
        ci.pack(fill="x", padx=14, pady=10)
        self._chk_box = tk.Label(ci, text="☐", font=("Segoe UI", 13), bg=ACCENT_L, fg=ACCENT, cursor="hand2")
        self._chk_box.pack(side="left")
        ts = tk.Frame(ci, bg=ACCENT_L)
        ts.pack(side="left", padx=10)
        tk.Label(
            ts,
            text="Responsável financeiro é o próprio aluno / mesmo que o Contratante",
            font=F10B, bg=ACCENT_L, fg=ACCENT,
        ).pack(anchor="w")
        tk.Label(
            ts,
            text="Os dados do Contratante serão copiados automaticamente ao gerar o documento.",
            font=F9, bg=ACCENT_L, fg=TEXT2,
        ).pack(anchor="w")
        for w in [toggle, ci, self._chk_box, ts] + list(ts.winfo_children()):
            w.bind("<Button-1>", self._toggle_mesmo)

        s1 = Section(f, "Dados do Responsável Financeiro", "💳")
        s1.pack(fill="x", padx=20, pady=6)
        g1 = build_fields(s1, [
            ("Nome completo",      "resp_nome",        30, "text"),
            ("Grau de Parentesco", "resp_parent",      18, "text"),
            ("Estado Civil",       "resp_estado_civil",14, "text"),
            ("Data de Nasc.",      "resp_nasc",        14, "text", "DD/MM/AAAA"),
            ("RG",                 "resp_rg",          18, "text"),
            ("CPF",                "resp_cpf",         18, "cpf",  "000.000.000-00"),
        ], self.v, self)

        s2 = Section(f, "Endereço do Responsável — CEP automático", "📍")
        s2.pack(fill="x", padx=20, pady=6)
        g2 = build_fields(s2, [
            ("CEP",         "resp_cep",    10, "cep_resp", "00000-000"),
            ("Número",      "resp_numero",  8, "text"),
            ("Logradouro",  "resp_end",    28, "text"),
            ("Complemento", "resp_comp",   18, "text"),
            ("Bairro",      "resp_bairro", 20, "text"),
            ("Cidade",      "resp_cidade", 20, "text"),
            ("UF",          "resp_uf",      5, "text"),
        ], self.v, self)

        s3 = Section(f, "Contato do Responsável", "📱")
        s3.pack(fill="x", padx=20, pady=(6, 20))
        g3 = build_fields(s3, [
            ("Telefone", "resp_tel",   16, "fone", "(XX) XXXXX-XXXX"),
            ("E-mail",   "resp_email", 28, "text"),
        ], self.v, self)

        self._resp_entries = []
        for grid in (g1, g2, g3):
            for child in grid.winfo_children():
                if isinstance(child, LabeledEntry):
                    self._resp_entries.append(child)

        return sc

    def _toggle_mesmo(self, _=None):
        val = not self.v["_mesmo"].get()
        self.v["_mesmo"].set(val)
        self._chk_box.configure(text="☑" if val else "☐")
        state = "disabled" if val else "normal"
        for entry in self._resp_entries:
            try:
                entry.entry.configure(state=state)
            except tk.TclError:
                pass

    def _pg_cursos(self):
        sc = ScrollableFrame(self._body)
        f = sc.inner

        s1 = Section(f, "Adicionar / Editar Curso", "✏")
        s1.pack(fill="x", padx=20, pady=(16, 6))
        form = tk.Frame(s1, bg=CARD)
        form.pack(fill="x", padx=20, pady=(8, 4))
        form.columnconfigure(0, weight=2)
        form.columnconfigure(1, weight=2)

        tk.Label(form, text="Nome do Curso", font=F9B, bg=CARD, fg=TEXT2).grid(
            row=0, column=0, sticky="w", pady=(0, 3)
        )
        self._nc_nome = tk.StringVar()
        name_border = tk.Frame(form, bg=BORDER2)
        name_border.grid(row=1, column=0, sticky="ew", padx=(0, 12), pady=(0, 10))
        name_inner = tk.Frame(name_border, bg=WHITE)
        name_inner.pack(fill="x", padx=1, pady=1)
        self._nc_nome_e = tk.Entry(
            name_inner, textvariable=self._nc_nome, font=F10, bg=WHITE, fg=TEXT,
            bd=0, relief="flat", width=30, insertbackground=ACCENT,
        )
        self._nc_nome_e.pack(padx=10, pady=7, fill="x")
        self._nc_nome_e.bind("<FocusIn>",  lambda e: name_border.configure(bg=ACCENT))
        self._nc_nome_e.bind("<FocusOut>", lambda e: name_border.configure(bg=BORDER2))

        tk.Label(form, text="Horário Padrão", font=F9B, bg=CARD, fg=TEXT2).grid(
            row=0, column=1, sticky="w", pady=(0, 3)
        )
        self._nc_hora = tk.StringVar()
        hora_border = tk.Frame(form, bg=BORDER2)
        hora_border.grid(row=1, column=1, sticky="ew", pady=(0, 10))
        hora_inner = tk.Frame(hora_border, bg=WHITE)
        hora_inner.pack(fill="x", padx=1, pady=1)
        self._nc_hora_e = tk.Entry(
            hora_inner, textvariable=self._nc_hora, font=F10, bg=WHITE, fg=TEXT,
            bd=0, relief="flat", width=28, insertbackground=ACCENT,
        )
        self._nc_hora_e.pack(padx=10, pady=7, fill="x")
        self._nc_hora_e.bind("<FocusIn>",  lambda e: hora_border.configure(bg=ACCENT))
        self._nc_hora_e.bind("<FocusOut>", lambda e: hora_border.configure(bg=BORDER2))

        btn_row = tk.Frame(s1, bg=CARD)
        btn_row.pack(anchor="w", padx=20, pady=(0, 14))
        Btn(btn_row, "✅ Salvar Curso",       self._salvar_curso,  SUCCESS).pack(side="left", padx=(0, 8))
        Btn(btn_row, "✏ Editar Selecionado", self._editar_curso,  ACCENT, light=True, small=True).pack(side="left", padx=(0, 8))
        Btn(btn_row, "🗑 Remover",            self._remover_curso, DANGER, small=True).pack(side="left")

        s2 = Section(f, "Cursos Cadastrados", "📋")
        s2.pack(fill="x", padx=20, pady=(6, 6))
        list_frame = tk.Frame(s2, bg=CARD)
        list_frame.pack(fill="x", padx=20, pady=(8, 16))
        header = tk.Frame(list_frame, bg=ACCENT_L)
        header.pack(fill="x")
        tk.Label(header, text="  Curso", font=F9B, bg=ACCENT_L, fg=ACCENT, width=38, anchor="w").pack(
            side="left", pady=6, padx=4
        )
        tk.Label(header, text="Horário Padrão", font=F9B, bg=ACCENT_L, fg=ACCENT, anchor="w").pack(
            side="left", pady=6
        )
        lb_frame = tk.Frame(list_frame, bg=BORDER2)
        lb_frame.pack(fill="x", pady=(1, 0))
        self._cursos_listbox = tk.Listbox(
            lb_frame, font=F10, bg=WHITE, fg=TEXT,
            selectbackground=ACCENT, selectforeground=WHITE,
            bd=0, relief="flat", height=10, activestyle="none", highlightthickness=0,
        )
        list_sb = tk.Scrollbar(lb_frame, command=self._cursos_listbox.yview)
        self._cursos_listbox.configure(yscrollcommand=list_sb.set)
        self._cursos_listbox.pack(side="left", fill="both", expand=True)
        list_sb.pack(side="right", fill="y")
        self._cursos_listbox.bind("<<ListboxSelect>>", self._on_curso_select)

        s3 = Section(f, "Importar de Planilha Excel", "📂")
        s3.pack(fill="x", padx=20, pady=(6, 20))
        import_frame = tk.Frame(s3, bg=CARD)
        import_frame.pack(fill="x", padx=20, pady=(8, 16))
        tk.Label(
            import_frame,
            text="Coluna A = nome do curso   |   Coluna B = horário (opcional)",
            font=F9, bg=CARD, fg=TEXT3,
        ).pack(anchor="w", pady=(0, 10))
        Btn(import_frame, "📂 Importar Planilha .xlsx", self._importar_planilha, ACCENT, light=True).pack(
            anchor="w"
        )

        self._refresh_lista_cursos()
        return sc

    def _salvar_curso(self):
        nome = self._nc_nome.get().strip()
        hora = self._nc_hora.get().strip()
        if not nome:
            messagebox.showwarning("Atenção", "Digite o nome do curso.")
            return
        for i, curso in enumerate(self.cursos_lista):
            if curso["nome"].lower() == nome.lower():
                self.cursos_lista[i] = {"nome": nome, "horario": hora}
                self._save_cursos()
                self._refresh_lista_cursos()
                self._status(f"✓ '{nome}' atualizado!", SUCCESS)
                return
        self.cursos_lista.append({"nome": nome, "horario": hora})
        self._save_cursos()
        self._refresh_lista_cursos()
        self._nc_nome.set("")
        self._nc_hora.set("")
        self._status(f"✓ '{nome}' adicionado!", SUCCESS)

    def _editar_curso(self):
        sel = self._cursos_listbox.curselection()
        if not sel:
            messagebox.showwarning("Atenção", "Selecione um curso.")
            return
        curso = self.cursos_lista[sel[0]]
        self._nc_nome.set(curso["nome"])
        self._nc_hora.set(curso["horario"])
        self._nc_nome_e.focus_set()

    def _remover_curso(self):
        sel = self._cursos_listbox.curselection()
        if not sel:
            messagebox.showwarning("Atenção", "Selecione um curso.")
            return
        nome = self.cursos_lista[sel[0]]["nome"]
        if messagebox.askyesno("Remover", f"Remover '{nome}'?"):
            self.cursos_lista.pop(sel[0])
            self._save_cursos()
            self._refresh_lista_cursos()
            self._status(f"🗑 '{nome}' removido.", TEXT2)

    def _on_curso_select(self, _=None):
        sel = self._cursos_listbox.curselection()
        if not sel:
            return
        curso = self.cursos_lista[sel[0]]
        self.v["curso"].set(curso["nome"])
        if curso["horario"]:
            self.v["dias_horario"].set(curso["horario"])

    def _on_combo_curso(self, _=None):
        nome = self.v["curso"].get().strip()
        for curso in self.cursos_lista:
            if curso["nome"] == nome:
                if curso["horario"]:
                    self.v["dias_horario"].set(curso["horario"])
                break

    def _refresh_lista_cursos(self):
        self._cursos_listbox.delete(0, "end")
        for i, curso in enumerate(self.cursos_lista):
            hora = curso["horario"] or "—"
            self._cursos_listbox.insert("end", f"  {curso['nome']:<38}{hora}")
            self._cursos_listbox.itemconfigure(i, background=WHITE if i % 2 == 0 else "#F8FAFC")
        self._curso_combo["values"] = [c["nome"] for c in self.cursos_lista]

    def _save_cursos(self):
        try:
            cfg = {}
            if os.path.exists(DRAFT_FILE):
                with open(DRAFT_FILE, encoding="utf-8") as fp:
                    cfg = json.load(fp)
            cfg["_cursos"] = self.cursos_lista
            with open(DRAFT_FILE, "w", encoding="utf-8") as fp:
                json.dump(cfg, fp, ensure_ascii=False, indent=2)
        except Exception as e:
            messagebox.showerror("Erro", str(e))

    def _importar_planilha(self):
        if not HAS_OPENPYXL:
            messagebox.showwarning("Atenção", "pip install openpyxl")
            return
        path = filedialog.askopenfilename(
            title="Planilha",
            filetypes=[("Excel", "*.xlsx *.xls"), ("Todos", "*.*")],
        )
        if not path:
            return
        try:
            import openpyxl
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            ws = wb.active
            imported = 0
            for row in ws.iter_rows(values_only=True):
                nome = str(row[0]).strip() if row[0] else ""
                hora = str(row[1]).strip() if len(row) > 1 and row[1] else ""
                if not nome or nome.lower() in ("curso", "nome", ""):
                    continue
                if not any(c["nome"].lower() == nome.lower() for c in self.cursos_lista):
                    self.cursos_lista.append({"nome": nome, "horario": hora})
                    imported += 1
            wb.close()
            self._save_cursos()
            self._refresh_lista_cursos()
            self._status(f"✓ {imported} cursos importados!", SUCCESS)
        except Exception as e:
            messagebox.showerror("Erro", str(e))

    def _pg_config(self):
        sc = ScrollableFrame(self._body)
        f = sc.inner

        s1 = Section(f, "Modelo de Contrato", "📄")
        s1.pack(fill="x", padx=20, pady=(16, 6))
        mf = tk.Frame(s1, bg=CARD)
        mf.pack(fill="x", padx=20, pady=(8, 16))
        tk.Label(
            mf,
            text="Selecionado uma vez — não precisa escolher toda vez que gerar.",
            font=F9, bg=CARD, fg=TEXT3,
        ).pack(anchor="w", pady=(0, 8))
        path_box = tk.Frame(mf, bg=ACCENT_L, highlightbackground=BORDER, highlightthickness=1)
        path_box.pack(fill="x", pady=(0, 10))
        self._modelo_lbl = tk.Label(
            path_box, textvariable=self.v["_modelo_path"],
            font=F9, bg=ACCENT_L, fg=TEXT2, anchor="w", pady=8, padx=12, wraplength=560,
        )
        self._modelo_lbl.pack(fill="x")
        btn_row = tk.Frame(mf, bg=CARD)
        btn_row.pack(anchor="w")
        Btn(btn_row, "📂 Selecionar Modelo", self._selecionar_modelo, ACCENT).pack(side="left")
        Btn(btn_row, "✖ Remover", self._remover_modelo, "#94A3B8", small=True).pack(
            side="left", padx=10
        )

        s2 = Section(f, "Dados da Escola", "🏫")
        s2.pack(fill="x", padx=20, pady=6)
        build_fields(s2, [
            ("Nome da Escola",  "escola_nome",     32, "text"),
            ("CNPJ",            "escola_cnpj",     20, "text"),
            ("Endereço",        "escola_end",      32, "text"),
            ("Município/UF",    "escola_municipio",20, "text"),
            ("Telefone",        "escola_tel",      18, "text"),
            ("E-mail",          "escola_email",    26, "text"),
            ("Diretor(a)",      "escola_diretor",  26, "text"),
            ("Coordenador(a)",  "escola_coord",    26, "text"),
        ], self.v, self)

        s3 = Section(f, "Dicas de Uso", "💡")
        s3.pack(fill="x", padx=20, pady=(6, 20))
        tips = [
            ("🔍 CEP automático",    "Endereço preenchido via ViaCEP."),
            ("✅ CPF validado",       "Dígito verificador em tempo real."),
            ("📱 Telefone formatado", "(XX) XXXXX-XXXX ao digitar."),
            ("📋 Cursos salvos",     "Cadastre em 'Gerenciar Cursos'."),
            ("📄 Modelo fixo",       "Configure acima — não precisa selecionar toda vez."),
            ("💾 Rascunho",          "Use Salvar para retomar depois."),
            ("⚡ Resp. = Aluno",      "Checkbox copia os dados do contratante."),
        ]
        tips_grid = tk.Frame(s3, bg=CARD)
        tips_grid.pack(fill="x", padx=20, pady=(8, 16))
        for i, (title, desc) in enumerate(tips):
            row_bg = "#F8FAFC" if i % 2 == 0 else CARD
            row = tk.Frame(tips_grid, bg=row_bg)
            row.pack(fill="x", pady=1)
            tk.Label(row, text=title, font=F9B, bg=row_bg, fg=TEXT, width=22, anchor="w").pack(
                side="left", padx=12, pady=8
            )
            tk.Label(row, text=desc, font=F9, bg=row_bg, fg=TEXT2, anchor="w").pack(
                side="left", padx=4
            )

        return sc

    def _pg_placeholders(self):
        sc = ScrollableFrame(self._body)
        s = Section(sc.inner, "Placeholders para o Modelo Word", "📌")
        s.pack(fill="x", padx=20, pady=(16, 20))

        code_frame = tk.Frame(s, bg="#1E293B")
        code_frame.pack(fill="x", padx=20, pady=(8, 4))
        txt = tk.Text(
            code_frame, height=32, font=("Courier New", 9),
            bg="#1E293B", fg="#94A3B8", bd=0, relief="flat", wrap="none",
            insertbackground=ACCENT,
        )
        txt.pack(fill="x", padx=12, pady=12)
        txt.insert("end", GUIDE)
        txt.configure(state="disabled")

        btn_row = tk.Frame(s, bg=CARD)
        btn_row.pack(anchor="e", padx=20, pady=(4, 16))
        Btn(
            btn_row,
            "📋 Copiar tudo",
            lambda: [
                self.clipboard_clear(),
                self.clipboard_append(GUIDE),
                messagebox.showinfo("✅", "Copiado!"),
            ],
            ACCENT, light=True, small=True,
        ).pack()

        return sc

    def _lookup_cep(self, cep, target):
        if not HAS_REQUESTS:
            self._status("⚠ instale requests", WARN)
            return
        self._status("🔍 Buscando CEP…", ACCENT)
        threading.Thread(
            target=lambda: self.after(0, lambda: self._fill_cep(buscar_cep(cep), target)),
            daemon=True,
        ).start()

    def _fill_cep(self, data, target):
        if not data:
            self._status("❌ CEP não encontrado", DANGER)
            return
        field_map = (
            {"logradouro": "endereco", "bairro": "bairro", "localidade": "cidade", "uf": "uf"}
            if target == "cont"
            else {"logradouro": "resp_end", "bairro": "resp_bairro", "localidade": "resp_cidade", "uf": "resp_uf"}
        )
        for api_key, var_key in field_map.items():
            if var_key in self.v and data.get(api_key):
                self.v[var_key].set(data[api_key])
        self._status("✓ Endereço preenchido!", SUCCESS)

    def _selecionar_modelo(self):
        path = filedialog.askopenfilename(
            title="Modelo (.docx)", filetypes=[("Word", "*.docx"), ("Todos", "*.*")]
        )
        if path:
            self.v["_modelo_path"].set(path)
            self._save_draft()
            self._status("✓ Modelo salvo!", SUCCESS)

    def _remover_modelo(self):
        self.v["_modelo_path"].set("")
        self._save_draft()
        self._status("Modelo removido.", TEXT2)

    def _get_template_path(self):
        saved = self.v.get("_modelo_path", tk.StringVar()).get()
        if saved and os.path.exists(saved):
            return saved
        if saved and not os.path.exists(saved):
            messagebox.showwarning(
                "Modelo não encontrado",
                f"Arquivo não encontrado:\n{saved}\nSelecione outro.",
            )
        path = filedialog.askopenfilename(
            title="Modelo (.docx)", filetypes=[("Word", "*.docx"), ("Todos", "*.*")]
        )
        if path:
            self.v["_modelo_path"].set(path)
            self._save_draft()
        return path or None

    def _choose_save_path(self, ext):
        nome = self.v.get("nome", tk.StringVar()).get().strip().replace(" ", "_") or "contrato"
        return (
            filedialog.asksaveasfilename(
                title="Salvar como…",
                defaultextension=f".{ext}",
                initialfile=f"Contrato_{nome}.{ext}",
                filetypes=[(ext.upper(), f"*.{ext}")],
            )
            or None
        )

    def _save_draft(self):
        data = {k: v.get() for k, v in self.v.items() if not k.startswith("_")}
        data["_obs"] = self._obs.get("1.0", "end").strip()
        data["_cursos"] = self.cursos_lista
        modelo = self.v.get("_modelo_path", tk.StringVar()).get()
        if modelo:
            data["_modelo_path"] = modelo
        try:
            with open(DRAFT_FILE, "w", encoding="utf-8") as fp:
                json.dump(data, fp, ensure_ascii=False, indent=2)
            self._status("💾 Salvo!", SUCCESS)
        except Exception as e:
            messagebox.showerror("Erro", str(e))

    def _load_draft(self):
        if not os.path.exists(DRAFT_FILE):
            return
        try:
            with open(DRAFT_FILE, encoding="utf-8") as fp:
                data = json.load(fp)
            for key, val in data.items():
                if key == "_obs":
                    self._obs.delete("1.0", "end")
                    self._obs.insert("end", val)
                elif key == "_cursos" and isinstance(val, list):
                    self.cursos_lista = val
                    self._refresh_lista_cursos()
                elif key == "_modelo_path" and val:
                    self.v["_modelo_path"].set(val)
                elif key in self.v:
                    self.v[key].set(val)
            self._status("📂 Rascunho carregado", ACCENT)
        except Exception:
            pass

    def _confirm_clear(self):
        if messagebox.askyesno("Limpar", "Deseja limpar todos os campos?"):
            for key, var in self.v.items():
                if not key.startswith("_"):
                    var.set("")
            self._obs.delete("1.0", "end")
            self._status("🗑 Limpo", TEXT2)

    def _build_mapping(self):
        def get(key):
            return self.v[key].get().strip() if key in self.v else ""

        if self.v.get("_mesmo") and self.v["_mesmo"].get():
            copies = [
                ("nome",        "resp_nome"),
                ("rg",          "resp_rg"),
                ("cpf",         "resp_cpf"),
                ("endereco",    "resp_end"),
                ("bairro",      "resp_bairro"),
                ("cidade",      "resp_cidade"),
                ("uf",          "resp_uf"),
                ("cep",         "resp_cep"),
                ("telefone1",   "resp_tel"),
                ("data_nasc",   "resp_nasc"),
                ("numero",      "resp_numero"),
                ("complemento", "resp_comp"),
                ("email",       "resp_email"),
            ]
            for src, dst in copies:
                self.v[dst].set(get(src))
            self.v["resp_parent"].set("Próprio aluno")

        return {
            "{{NOME}}":              get("nome"),
            "{{NATURALIDADE}}":      get("naturalidade"),
            "{{DATA_NASC}}":         get("data_nasc"),
            "{{RG}}":                get("rg"),
            "{{CPF}}":               get("cpf"),
            "{{CEP}}":               get("cep"),
            "{{ENDERECO}}":          get("endereco"),
            "{{NUMERO}}":            get("numero"),
            "{{COMPLEMENTO}}":       get("complemento"),
            "{{BAIRRO}}":            get("bairro"),
            "{{CIDADE}}":            get("cidade"),
            "{{UF}}":                get("uf"),
            "{{TELEFONE1}}":         get("telefone1"),
            "{{TELEFONE2}}":         get("telefone2"),
            "{{TELEFONES}}":         " / ".join(filter(None, [get("telefone1"), get("telefone2")])),
            "{{EMAIL}}":             get("email"),
            "{{CURSO}}":             get("curso"),
            "{{DIAS_HORARIO}}":      get("dias_horario"),
            "{{INICIO_1}}":          get("inicio_1"),
            "{{INICIO_2}}":          get("inicio_2"),
            "{{INICIO_3}}":          get("inicio_3"),
            "{{VALOR_TOTAL}}":       fmt_cur(get("valor_total")),
            "{{PARCELA_NORMAL}}":    fmt_cur(get("parcela_normal")),
            "{{PARCELA_DESC}}":      fmt_cur(get("parcela_desc")),
            "{{DESCONTO_PCT}}":      get("desconto_pct") + "%",
            "{{QTD_PARCELAS}}":      get("qtd_parcelas"),
            "{{DIA_VENCIMENTO}}":    get("dia_vencimento"),
            "{{VALOR_MATRICULA}}":   fmt_cur(get("valor_matricula")),
            "{{DATA_1PAG}}":         get("data_1pag"),
            "{{MES_REF}}":           get("mes_ref"),
            "{{DATA_VENC_SEM_DESC}}":get("data_venc_sem_desc"),
            "{{FORMA_PAGAMENTO}}":   get("forma_pagamento"),
            "{{OBS}}":               self._obs.get("1.0", "end").strip(),
            "{{ESCOLA_NOME}}":       get("escola_nome"),
            "{{ESCOLA_CNPJ}}":       get("escola_cnpj"),
            "{{ESCOLA_END}}":        get("escola_end"),
            "{{ESCOLA_MUNICIPIO}}":  get("escola_municipio"),
            "{{ESCOLA_TEL}}":        get("escola_tel"),
            "{{ESCOLA_EMAIL}}":      get("escola_email"),
            "{{ESCOLA_DIRETOR}}":    get("escola_diretor"),
            "{{ESCOLA_COORD}}":      get("escola_coord"),
            "{{RESP_NOME}}":         get("resp_nome"),
            "{{RESP_PARENT}}":       get("resp_parent"),
            "{{RESP_ESTADO_CIVIL}}": get("resp_estado_civil"),
            "{{RESP_NASC}}":         get("resp_nasc"),
            "{{RESP_RG}}":           get("resp_rg"),
            "{{RESP_CPF}}":          get("resp_cpf"),
            "{{RESP_CEP}}":          get("resp_cep"),
            "{{RESP_END}}":          get("resp_end"),
            "{{RESP_NUMERO}}":       get("resp_numero"),
            "{{RESP_COMP}}":         get("resp_comp"),
            "{{RESP_BAIRRO}}":       get("resp_bairro"),
            "{{RESP_CIDADE}}":       get("resp_cidade"),
            "{{RESP_UF}}":           get("resp_uf"),
            "{{RESP_TEL}}":          get("resp_tel"),
            "{{RESP_EMAIL}}":        get("resp_email"),
        }

    def gerar(self, modo):
        template = self._get_template_path()
        if not template:
            return
        mapping = self._build_mapping()

        if modo in ("word", "ambos"):
            out = self._choose_save_path("docx")
            if out:
                try:
                    doc = Document(template)
                    replace_all(doc, mapping)
                    doc.save(out)
                    if modo == "word":
                        messagebox.showinfo("✅ Sucesso", f"Word gerado!\n{out}")
                except Exception as e:
                    messagebox.showerror("Erro", str(e))
                    return

        if modo in ("pdf", "ambos"):
            out_pdf = self._choose_save_path("pdf")
            if out_pdf:
                tmp = out_pdf.replace(".pdf", "_tmp.docx")
                try:
                    doc = Document(template)
                    replace_all(doc, mapping)
                    doc.save(tmp)
                    ok, err = convert_to_pdf(tmp, out_pdf)
                    try:
                        os.remove(tmp)
                    except OSError:
                        pass
                    if ok:
                        messagebox.showinfo("✅ Sucesso", f"PDF gerado!\n{out_pdf}")
                    else:
                        messagebox.showwarning("PDF não gerado", err)
                except Exception as e:
                    messagebox.showerror("Erro", str(e))

        if modo == "ambos":
            messagebox.showinfo("✅ Concluído", "Word e PDF gerados!")
